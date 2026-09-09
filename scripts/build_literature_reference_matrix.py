#!/usr/bin/env python3
"""Build literature review Excel matrix from reference.txt."""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "docs" / "literature_review_reference_matrix.xlsx"

ROWS = [
    dict(
        ref_no=1,
        authors="Almuhanna & Dardouri",
        year=2025,
        title="A deep learning/machine learning approach for anomaly-based network intrusion detection",
        venue="Frontiers in Artificial Intelligence",
        doi_or_id="10.3389/frai.2025.1625891",
        category="Security anomaly / NIDS",
        method_summary="DL/ML pipeline for anomaly-based network intrusion detection.",
        key_contribution="Shows modern ML/DL can detect network intrusions from traffic features.",
        gaps_limitations="Focused on network intrusion, not container CPU forecasting; little on long-term resource planning or hybrid Prophet+GRU lifecycle.",
        maps_to_our_gap="Supports Module 4 vision (security AD); limited Module 1 relevance.",
        relevance_to_fyp="Medium (Module 4)",
        related_module="Module 4",
        our_response="Cite as security-anomaly related work; not a forecasting baseline.",
    ),
    dict(
        ref_no=2,
        authors="Babaey & Faragardi",
        year=2025,
        title="Detecting zero-day web attacks with an ensemble of LSTM, GRU, and stacked autoencoders",
        venue="Computers",
        doi_or_id="10.3390/computers14060205",
        category="Security anomaly / web attacks",
        method_summary="Ensemble of LSTM, GRU, and stacked autoencoders for zero-day web attack detection.",
        key_contribution="Demonstrates sequence models (LSTM/GRU) + AE for security anomaly detection.",
        gaps_limitations="Security-domain focus; no cloud resource forecasting; no deployment lifecycle for forecasting models.",
        maps_to_our_gap="Motivates GRU usefulness and Module 4; not Module 1 core comparison.",
        relevance_to_fyp="Medium (Module 4)",
        related_module="Module 4",
        our_response="Cite for GRU/AE in anomaly detection; separate from Module 1 forecasting claims.",
    ),
    dict(
        ref_no=3,
        authors="Bountzis et al.",
        year=2025,
        title="A deep one-class classifier for network anomaly detection using autoencoders and one-class SVM",
        venue="Frontiers in Computer Science",
        doi_or_id="10.3389/fcomp.2025.1646679",
        category="Security / network anomaly",
        method_summary="Autoencoder + one-class SVM for network anomaly detection.",
        key_contribution="One-class deep approach for anomaly detection when attack labels are scarce.",
        gaps_limitations="Not resource forecasting; no container workload modeling; no concept-drift lifecycle for forecasters.",
        maps_to_our_gap="Module 3/4 related work.",
        relevance_to_fyp="Medium (Module 3/4)",
        related_module="Module 3/4",
        our_response="Cite under anomaly-detection literature for later modules.",
    ),
    dict(
        ref_no=4,
        authors="Cho et al.",
        year=2014,
        title="Learning phrase representations using RNN encoder-decoder for statistical machine translation",
        venue="EMNLP",
        doi_or_id="EMNLP 2014, pp. 1724-1734",
        category="Foundation — GRU",
        method_summary="Introduces GRU as a gated RNN encoder-decoder component.",
        key_contribution="Foundational GRU architecture widely used in sequence modeling.",
        gaps_limitations="NLP/MT paper; not cloud forecasting; no Prophet hybrid or drift lifecycle.",
        maps_to_our_gap="Foundation for Architecture 1 & 2 (GRU residual / Global GRU).",
        relevance_to_fyp="High (foundation)",
        related_module="Module 1",
        our_response="Cite as GRU foundation; we apply GRU to residual/global CPU forecasting.",
    ),
    dict(
        ref_no=5,
        authors="Correia et al.",
        year=2024,
        title="Online model-based anomaly detection in multivariate time series: taxonomy, survey, research challenges and future directions",
        venue="arXiv",
        doi_or_id="arXiv:2408.03747",
        category="Survey — online anomaly / time series",
        method_summary="Taxonomy and survey of online model-based multivariate TS anomaly detection.",
        key_contribution="Maps challenges in online detection, drift, and streaming evaluation.",
        gaps_limitations="Survey of anomaly detection, not forecasting accuracy; no Hybrid Prophet+GRU for containers.",
        maps_to_our_gap="Gap 5 (adaptive/online); informs AFMLF motivation indirectly.",
        relevance_to_fyp="Medium",
        related_module="Module 1/3",
        our_response="Cite for online/adaptive challenges; AFMLF addresses lifecycle for forecasting, not a new AD algorithm.",
    ),
    dict(
        ref_no=6,
        authors="Dkmak et al.",
        year=2025,
        title="AI-driven anomaly detection in cloud-native microservices: the Night's Watch algorithm",
        venue="Applied Sciences",
        doi_or_id="10.3390/app152312762",
        category="Cloud-native anomaly detection",
        method_summary="AI anomaly detection tailored to cloud-native microservices.",
        key_contribution="Cloud-native microservice anomaly focus aligned with containerized environments.",
        gaps_limitations="Detection-centric; not long-term CPU forecasting; no hybrid statistical+DL forecaster comparison.",
        maps_to_our_gap="Module 3; environmental relevance to containers.",
        relevance_to_fyp="Medium (Module 3)",
        related_module="Module 3",
        our_response="Cite as cloud-native AD related work; Module 1 remains forecasting.",
    ),
    dict(
        ref_no=7,
        authors="Dogani, Khunjush & Seydali",
        year=2023,
        title="Host load prediction in cloud computing with DWT and BiGRU",
        venue="Computer Communications",
        doi_or_id="10.1016/j.comcom.2022.11.018",
        category="Cloud load forecasting",
        method_summary="Discrete wavelet transform + BiGRU for host load prediction.",
        key_contribution="Strong DL host-load forecasting baseline using BiGRU and multi-resolution features.",
        gaps_limitations="Host-level (not container cohort); wavelet+BiGRU only (no Prophet hybrid); limited unseen-container / lifecycle study.",
        maps_to_our_gap="G1, G2 — closest pure DL cloud-load baselines.",
        relevance_to_fyp="High",
        related_module="Module 1",
        our_response="Compare conceptually: we use Prophet+GRU residual hybrid and Global GRU under locked Day-1 protocol on Alibaba containers.",
    ),
    dict(
        ref_no=8,
        authors="El Mahdaouy et al.",
        year=2026,
        title="Deep learning for contextualized NetFlow-based network intrusion detection",
        venue="arXiv",
        doi_or_id="arXiv:2602.05594",
        category="Security / NIDS",
        method_summary="DL methods for NetFlow-based intrusion detection with evaluation/deployment discussion.",
        key_contribution="Brings deployment considerations into NetFlow NIDS research.",
        gaps_limitations="Security telemetry, not CPU forecasting; different data modality.",
        maps_to_our_gap="Module 4 related work.",
        relevance_to_fyp="Low–Medium (Module 4)",
        related_module="Module 4",
        our_response="Cite for security module literature; not Module 1 baseline.",
    ),
    dict(
        ref_no=9,
        authors="Gama et al.",
        year=2014,
        title="A survey on concept drift adaptation",
        venue="ACM Computing Surveys",
        doi_or_id="10.1145/2523813",
        category="Survey — concept drift",
        method_summary="Seminal survey of concept drift types and adaptation strategies.",
        key_contribution="Defines drift taxonomy and adaptation families used across streaming ML.",
        gaps_limitations="General survey; not a container forecasting system; no Hybrid Prophet+GRU lifecycle design.",
        maps_to_our_gap="G5 — primary drift literature anchor for AFMLF.",
        relevance_to_fyp="High",
        related_module="Module 1 (AFMLF)",
        our_response="Cite as drift foundation; AFMLF is lifecycle governance for a frozen Hybrid, not a new drift algorithm.",
    ),
    dict(
        ref_no=10,
        authors="Ghatage, Patil & Shinde",
        year=2023,
        title="Lightweight RNN-based model for adaptive time series forecasting with concept drift detection in smart homes",
        venue="JESA",
        doi_or_id="10.18280/jesa.560609",
        category="Adaptive forecasting + drift",
        method_summary="Lightweight RNN forecasting with drift detection in smart-home time series.",
        key_contribution="Shows forecasting + drift detection can be combined in adaptive pipelines.",
        gaps_limitations="Smart-home domain; lightweight RNN only; no Prophet hybrid, no cohort retrain policy for global cloud models.",
        maps_to_our_gap="G5 — related adaptive forecasting.",
        relevance_to_fyp="High",
        related_module="Module 1 (AFMLF)",
        our_response="Closest adaptive-forecasting relatives; we specialize to container Hybrid + cohort AFMLF.",
    ),
    dict(
        ref_no=11,
        authors="Guruge & Priyadarshana",
        year=2025,
        title="Time series forecasting-based Kubernetes autoscaling using Facebook Prophet and LSTM",
        venue="Frontiers in Computer Science",
        doi_or_id="10.3389/fcomp.2025.1509165",
        category="Hybrid Prophet + DL (K8s)",
        method_summary="Prophet + LSTM forecasting to drive Kubernetes autoscaling.",
        key_contribution="Closest hybrid related work: Prophet + deep sequence model for container orchestration.",
        gaps_limitations="LSTM not GRU residual hybrid; autoscaling-centric metrics; no Global GRU comparison; no peak-aware/LFHE/AFMLF stack on Alibaba trace.",
        maps_to_our_gap="G2 — primary closest work for Hybrid novelty discussion.",
        relevance_to_fyp="Very High",
        related_module="Module 1",
        our_response="Differentiate: Prophet+GRU residual hybrid, Alibaba CPU Day-1 protocol, Global GRU comparison, residual/peak diagnostics, AFMLF lifecycle.",
    ),
    dict(
        ref_no=12,
        authors="Han et al.",
        year=2024,
        title="A hybrid loss framework for decomposition-based time series forecasting methods",
        venue="arXiv",
        doi_or_id="arXiv:2411.11340",
        category="Loss / decomposition forecasting",
        method_summary="Hybrid loss balancing global and component errors in decomposition-based TS forecasting.",
        key_contribution="Shows loss design matters for decomposition hybrids beyond architecture alone.",
        gaps_limitations="Not container-CPU specific; not DA-MSE residual variance-collapse test on Hybrid GRU.",
        maps_to_our_gap="Motivates LFHE / objective experiments.",
        relevance_to_fyp="High",
        related_module="Module 1 (LFHE)",
        our_response="Cite for loss-aware decomposition; LFHE tests DA-MSE vs MSE on CSRLE B0 residuals.",
    ),
    dict(
        ref_no=13,
        authors="Hochreiter & Schmidhuber",
        year=1997,
        title="Long short-term memory",
        venue="Neural Computation",
        doi_or_id="10.1162/neco.1997.9.8.1735",
        category="Foundation — LSTM",
        method_summary="Introduces LSTM gated recurrent architecture.",
        key_contribution="Foundational sequence model underlying many forecasting baselines.",
        gaps_limitations="Not a cloud system paper; no hybrid Prophet design.",
        maps_to_our_gap="Foundation for DL baselines and related Prophet+LSTM work.",
        relevance_to_fyp="High (foundation)",
        related_module="Module 1",
        our_response="Cite as LSTM foundation; we prefer GRU for efficiency in residual/global models.",
    ),
    dict(
        ref_no=14,
        authors="Jolliffe",
        year=2002,
        title="Principal Component Analysis (2nd ed.)",
        venue="Springer book",
        doi_or_id="Book reference",
        category="Foundation — PCA / dimensionality",
        method_summary="Canonical PCA reference for unsupervised dimensionality reduction.",
        key_contribution="Standard method for feature compression / variance analysis.",
        gaps_limitations="Not forecasting or container telemetry system.",
        maps_to_our_gap="Supporting method literature (possible preprocessing/anomaly features).",
        relevance_to_fyp="Low–Medium",
        related_module="Supporting",
        our_response="Cite if PCA used in anomaly/feature pipelines; not a Module 1 forecaster.",
    ),
    dict(
        ref_no=15,
        authors="Karim et al.",
        year=2021,
        title="BHyPreC: Bi-LSTM based hybrid RNN to predict CPU workload of cloud VM",
        venue="IEEE Access",
        doi_or_id="10.1109/ACCESS.2021.3113714",
        category="Cloud VM CPU forecasting",
        method_summary="Bi-LSTM hybrid recurrent model for VM CPU workload prediction.",
        key_contribution="Strong published hybrid DL baseline for cloud CPU workload.",
        gaps_limitations="VM-centric; Bi-LSTM hybrid (not Prophet+GRU); limited multi-container global comparison & lifecycle.",
        maps_to_our_gap="G1 & G2 — major existing cloud CPU forecasting work.",
        relevance_to_fyp="Very High",
        related_module="Module 1",
        our_response="Key comparator paper: we target containers, Prophet residual hybrid, Global GRU, locked evaluation + AFMLF.",
    ),
    dict(
        ref_no=16,
        authors="Kingma & Welling",
        year=2014,
        title="Auto-encoding variational Bayes",
        venue="ICLR",
        doi_or_id="ICLR 2014",
        category="Foundation — VAE",
        method_summary="Introduces variational autoencoders for generative latent modeling.",
        key_contribution="Foundation for many deep anomaly detectors using reconstruction/latent density.",
        gaps_limitations="Not forecasting; generative modeling paper.",
        maps_to_our_gap="Module 3/4 foundation for VAE-based AD.",
        relevance_to_fyp="Medium (Module 3/4)",
        related_module="Module 3/4",
        our_response="Cite as VAE foundation for anomaly modules.",
    ),
    dict(
        ref_no=17,
        authors="Liu, Ting & Zhou",
        year=2008,
        title="Isolation Forest",
        venue="IEEE ICDM",
        doi_or_id="10.1109/ICDM.2008.17",
        category="Foundation — anomaly detection",
        method_summary="Isolation Forest for unsupervised anomaly detection via random partitioning.",
        key_contribution="Widely used classical AD baseline.",
        gaps_limitations="Point/tabular AD; not long-term forecasting; no drift lifecycle for forecasters.",
        maps_to_our_gap="Module 3 foundation; also used in container extreme-load papers.",
        relevance_to_fyp="Medium (Module 3)",
        related_module="Module 3",
        our_response="Cite as classical AD baseline; Module 1 is forecasting.",
    ),
    dict(
        ref_no=18,
        authors="Liu, Zhang & Zhang",
        year=2025,
        title="A hybrid framework integrating traditional models and deep learning for multi-scale time series forecasting",
        venue="Entropy",
        doi_or_id="10.3390/e27070695",
        category="Hybrid statistical + DL forecasting",
        method_summary="Integrates traditional models with DL for multi-scale TS forecasting.",
        key_contribution="Supports the hybrid statistical+DL research direction.",
        gaps_limitations="General TS; not Alibaba container CPU; no AFMLF/Global GRU locked comparison.",
        maps_to_our_gap="G2 — hybrid paradigm support.",
        relevance_to_fyp="High",
        related_module="Module 1",
        our_response="Cite as hybrid-paradigm related work; our hybrid is Prophet+GRU residual for containers.",
    ),
    dict(
        ref_no=19,
        authors="Massey",
        year=1951,
        title="The Kolmogorov-Smirnov test for goodness of fit",
        venue="JASA",
        doi_or_id="10.1080/01621459.1951.10500769",
        category="Foundation — statistics",
        method_summary="Classical KS goodness-of-fit / distribution comparison test.",
        key_contribution="Statistical test often used in drift/distribution-shift detection.",
        gaps_limitations="Statistical tool only; not a forecasting system.",
        maps_to_our_gap="Supports drift/statistical testing methodology.",
        relevance_to_fyp="Medium (methods)",
        related_module="Module 1 (AFMLF methods)",
        our_response="Cite if KS/distribution tests used; AFMLF uses threshold+PH confirmation.",
    ),
    dict(
        ref_no=20,
        authors="Meng et al.",
        year=2026,
        title="BASE: burst-adaptive autoscaling via stacked ensembles for SLO assurance and cost efficiency",
        venue="IEEE Trans. Services Computing",
        doi_or_id="10.1109/TSC.2026.3668105",
        category="Autoscaling / systems",
        method_summary="Burst-adaptive autoscaling with stacked ensembles for SLO and cost.",
        key_contribution="Links prediction/adaptation to SLO and cost — deployment practicality.",
        gaps_limitations="Autoscaling system focus; not Prophet+GRU residual research stack; different evaluation goals.",
        maps_to_our_gap="G3 — deployment practicality / burst adaptation.",
        relevance_to_fyp="High",
        related_module="Module 1 (systems motivation)",
        our_response="Cite for why proactive forecasting matters for scaling; we contribute the forecasting+lifecycle layer.",
    ),
    dict(
        ref_no=21,
        authors="Nashaat et al.",
        year=2026,
        title="Dynamic machine learning approach for workload prediction in cloud environments",
        venue="Scientific Reports",
        doi_or_id="10.1038/s41598-026-40777-z",
        category="Cloud workload forecasting",
        method_summary="Dynamic ML approach for cloud workload prediction.",
        key_contribution="Recent cloud workload prediction emphasizing dynamic/adaptive ML.",
        gaps_limitations="May not provide Prophet+GRU residual hybrid + Global GRU fair protocol + AFMLF.",
        maps_to_our_gap="G1 & G5 — dynamic cloud prediction related work.",
        relevance_to_fyp="High",
        related_module="Module 1",
        our_response="Cite as recent dynamic cloud workload prediction; differentiate locked Hybrid/Global/AFMLF contributions.",
    ),
    dict(
        ref_no=22,
        authors="Nguyen et al.",
        year=2024,
        title="Variational autoencoder for anomaly detection: a comparative study",
        venue="arXiv",
        doi_or_id="arXiv:2408.13561",
        category="Survey/compare — VAE AD",
        method_summary="Comparative study of VAE variants for anomaly detection.",
        key_contribution="Benchmarks VAE-based AD design choices.",
        gaps_limitations="AD-focused; not resource forecasting.",
        maps_to_our_gap="Module 3 related work.",
        relevance_to_fyp="Medium (Module 3)",
        related_module="Module 3",
        our_response="Cite for VAE AD design choices in later modules.",
    ),
    dict(
        ref_no=23,
        authors="Nguyen & Watabe",
        year=2023,
        title="A method for network intrusion detection using flow sequence and BERT framework",
        venue="IEEE ICC",
        doi_or_id="10.1109/ICC45041.2023.10279335",
        category="Security / NIDS (transformers)",
        method_summary="Flow-sequence intrusion detection using BERT-style framework.",
        key_contribution="Shows transformers can model sequential network flows for IDS.",
        gaps_limitations="Security domain; heavy models; not container CPU forecasting.",
        maps_to_our_gap="Module 4; also answers 'why not transformers everywhere?'.",
        relevance_to_fyp="Low–Medium (Module 4)",
        related_module="Module 4",
        our_response="Cite for transformer/NIDS; Module 1 prioritizes interpretable hybrid + deployable GRU.",
    ),
    dict(
        ref_no=24,
        authors="Pawar et al.",
        year=2025,
        title="Enhancing forecasting accuracy in dynamic environments via PELT-driven drift detection and model adaptation",
        venue="Next Energy",
        doi_or_id="10.1016/j.nxener.2025.100462",
        category="Drift detection + forecast adaptation",
        method_summary="PELT change-point drift detection with model adaptation for dynamic forecasting.",
        key_contribution="Concrete drift-triggered adaptation pipeline for forecasting.",
        gaps_limitations="Energy-domain paper; PELT-based; not Hybrid Prophet+GRU container cohort lifecycle with diagnostics/human gate.",
        maps_to_our_gap="G5 — strong AFMLF related work.",
        relevance_to_fyp="Very High",
        related_module="Module 1 (AFMLF)",
        our_response="Differentiate: AFMLF uses per-container thresholds, PH confirm, Prophet-vs-Hybrid diagnostics, cohort trigger, offline candidate, deploy margin.",
    ),
    dict(
        ref_no=25,
        authors="Rossi et al.",
        year=2025,
        title="Forecasting workload in cloud computing: towards uncertainty-aware predictions and transfer learning",
        venue="Cluster Computing",
        doi_or_id="10.1007/s10586-024-04933-2",
        category="Cloud workload forecasting",
        method_summary="Cloud workload forecasting with uncertainty awareness and transfer learning.",
        key_contribution="Addresses generalization/transfer and predictive uncertainty in cloud forecasting.",
        gaps_limitations="Not the same Hybrid vs Global Prophet+GRU residual protocol; transfer ≠ our AFMLF lifecycle.",
        maps_to_our_gap="G1 & G3 — generalization / practical cloud forecasting.",
        relevance_to_fyp="Very High",
        related_module="Module 1",
        our_response="Cite for transfer/generalization; we contribute Global GRU + unseen demos + Hybrid baseline + AFMLF.",
    ),
    dict(
        ref_no=26,
        authors="Ryu et al.",
        year=2026,
        title="Hybrid runtime detection of malicious containers using eBPF",
        venue="CMC",
        doi_or_id="10.32604/cmc.2025.074871",
        category="Container security",
        method_summary="Hybrid runtime malicious-container detection using eBPF.",
        key_contribution="Runtime container security monitoring aligned with containerized deployment theme.",
        gaps_limitations="Security runtime, not CPU forecasting.",
        maps_to_our_gap="Module 4 / system vision.",
        relevance_to_fyp="Medium (Module 4)",
        related_module="Module 4",
        our_response="Cite for container security context of DracaSys.",
    ),
    dict(
        ref_no=27,
        authors="Taylor & Letham",
        year=2018,
        title="Forecasting at scale (Prophet)",
        venue="The American Statistician",
        doi_or_id="10.1080/00031305.2017.1380080",
        category="Foundation — Prophet",
        method_summary="Prophet additive model for trend/seasonality/holiday effects at scale.",
        key_contribution="Interpretable statistical forecasting backbone of our Hybrid architecture.",
        gaps_limitations="Weak on strong nonlinear residuals alone; not a deep residual learner; not multi-container global model.",
        maps_to_our_gap="G2 — motivates combining Prophet with GRU residuals.",
        relevance_to_fyp="Very High",
        related_module="Module 1",
        our_response="We use Prophet for trend/daily seasonality and GRU for residual correction.",
    ),
    dict(
        ref_no=28,
        authors="Torabi, Mirtaheri & Greco",
        year=2023,
        title="Practical autoencoder based anomaly detection by using vector reconstruction error",
        venue="Cybersecurity",
        doi_or_id="10.1186/s42400-022-00134-9",
        category="Anomaly detection — AE",
        method_summary="Practical AE anomaly detection using vector reconstruction error.",
        key_contribution="Practical guidance on AE reconstruction-error AD.",
        gaps_limitations="AD only; not forecasting.",
        maps_to_our_gap="Module 3 related work.",
        relevance_to_fyp="Medium (Module 3)",
        related_module="Module 3",
        our_response="Cite under AE anomaly detection literature.",
    ),
    dict(
        ref_no=29,
        authors="Vaswani et al.",
        year=2017,
        title="Attention is all you need",
        venue="NeurIPS",
        doi_or_id="NeurIPS 2017",
        category="Foundation — Transformer",
        method_summary="Introduces Transformer attention architecture.",
        key_contribution="Foundation for modern sequence/Transformer forecasting and NIDS models.",
        gaps_limitations="Not cloud hybrid forecasting; heavier than GRU for our Module 1 scope.",
        maps_to_our_gap="Answers 'why not Transformer?' examiner question.",
        relevance_to_fyp="Medium",
        related_module="Module 1 (scope choice)",
        our_response="Acknowledge Transformers; justify GRU/Prophet hybrid for interpretability, cost, and locked baselines.",
    ),
    dict(
        ref_no=30,
        authors="Wang & Wang",
        year=2024,
        title="Isolated forest-based prediction of container resource load extremes",
        venue="Applied Sciences",
        doi_or_id="10.3390/app14072911",
        category="Container load extremes / AD",
        method_summary="Isolation Forest for predicting/detecting container resource load extremes.",
        key_contribution="Container-specific extreme-load focus using Isolation Forest.",
        gaps_limitations="Extreme/anomaly oriented more than long-horizon forecasting accuracy; not Hybrid vs Global protocol.",
        maps_to_our_gap="G3 + peak/extreme motivation; Module 3 overlap.",
        relevance_to_fyp="High",
        related_module="Module 1/3",
        our_response="Cite for container extremes; our peak-aware track studies peak timesteps in forecasting loss.",
    ),
    dict(
        ref_no=31,
        authors="Wang et al.",
        year=2025,
        title="A survey of deep anomaly detection in multivariate time series",
        venue="Sensors",
        doi_or_id="10.3390/s25010190",
        category="Survey — deep TS anomaly",
        method_summary="Taxonomy/applications/directions for deep multivariate TS anomaly detection.",
        key_contribution="Broad map of deep TS AD methods and open challenges.",
        gaps_limitations="Survey of AD, not forecasting hybrids for CPU planning.",
        maps_to_our_gap="Module 3 literature backbone.",
        relevance_to_fyp="Medium (Module 3)",
        related_module="Module 3",
        our_response="Cite as AD survey for Modules 3/4.",
    ),
    dict(
        ref_no=32,
        authors="Xu et al.",
        year=2022,
        title="Anomaly Transformer: time series anomaly detection with association discrepancy",
        venue="ICLR",
        doi_or_id="ICLR 2022",
        category="Deep TS anomaly — Transformer",
        method_summary="Anomaly Transformer using association discrepancy for TS AD.",
        key_contribution="State-of-the-art style deep TS anomaly detector.",
        gaps_limitations="Detection not forecasting; different objective than Module 1.",
        maps_to_our_gap="Module 3 SOTA-related work.",
        relevance_to_fyp="Medium (Module 3)",
        related_module="Module 3",
        our_response="Cite as advanced TS AD; Module 1 contribution remains forecasting+lifecycle.",
    ),
    dict(
        ref_no=33,
        authors="Yu et al.",
        year=2024,
        title="LiNo: recursive residual decomposition of linear and nonlinear patterns for robust time series forecasting",
        venue="arXiv",
        doi_or_id="arXiv:2410.17159",
        category="Residual decomposition forecasting",
        method_summary="Recursive residual decomposition separating linear/nonlinear forecast patterns.",
        key_contribution="Supports residual-decomposition paradigm related to Hybrid forecasting.",
        gaps_limitations="General TS method; not Alibaba Hybrid Prophet+GRU + Global comparison + AFMLF.",
        maps_to_our_gap="G2 — residual decomposition related work.",
        relevance_to_fyp="High",
        related_module="Module 1",
        our_response="Cite for residual decomposition idea; our residual path is Prophet→GRU on container CPU.",
    ),
    dict(
        ref_no=34,
        authors="Yunita et al.",
        year=2025,
        title="Performance analysis of neural network architectures for time series forecasting: RNN, LSTM, GRU, and hybrid models",
        venue="MethodsX",
        doi_or_id="10.1016/j.mex.2025.103462",
        category="Architecture comparison (RNN family)",
        method_summary="Empirical comparison of RNN, LSTM, GRU, and hybrid neural forecasters.",
        key_contribution="Supports GRU as competitive efficient sequence model vs LSTM/RNN.",
        gaps_limitations="Generic architecture bake-off; not container Prophet hybrid or drift lifecycle.",
        maps_to_our_gap="Supports choosing GRU in Module 1.",
        relevance_to_fyp="High",
        related_module="Module 1",
        our_response="Cite to justify GRU selection vs LSTM for residual/global models.",
    ),
    dict(
        ref_no=35,
        authors="Zhang & Zhang",
        year=2025,
        title="Anomaly detection and objective security evaluation using AE, Isolation Forest, and multi-criteria decision methods",
        venue="Sensors",
        doi_or_id="10.3390/s25196250",
        category="Security evaluation / AD",
        method_summary="Combines AE, Isolation Forest, and MCDM for anomaly/security evaluation.",
        key_contribution="Multi-method AD + decision evaluation framing.",
        gaps_limitations="Security evaluation focus; not CPU forecasting.",
        maps_to_our_gap="Module 3/4 related work.",
        relevance_to_fyp="Medium (Module 3/4)",
        related_module="Module 3/4",
        our_response="Cite under combined AD/security evaluation literature.",
    ),
    dict(
        ref_no=36,
        authors="Zhang et al.",
        year=2023,
        title="OneNet: enhancing time series forecasting models under concept drift by online ensembling",
        venue="NeurIPS",
        doi_or_id="NeurIPS 2023",
        category="Concept drift — online forecasting",
        method_summary="Online ensembling to improve TS forecasting under concept drift.",
        key_contribution="Strong modern method for drift-aware forecasting via online ensembles.",
        gaps_limitations="Online ensembling algorithm; not a Hybrid Prophet+GRU container lifecycle with human-gated deploy.",
        maps_to_our_gap="G5 — major AFMLF related work.",
        relevance_to_fyp="Very High",
        related_module="Module 1 (AFMLF)",
        our_response="Differentiate carefully: we do not claim OneNet-style online ensembling; AFMLF is offline candidate lifecycle governance.",
    ),
    dict(
        ref_no=37,
        authors="Zong et al.",
        year=2018,
        title="Deep autoencoding Gaussian mixture model for unsupervised anomaly detection",
        venue="ICLR",
        doi_or_id="ICLR 2018",
        category="Foundation — deep AD (DAGMM)",
        method_summary="Combines deep autoencoder with GMM for unsupervised AD.",
        key_contribution="Influential deep unsupervised anomaly detection baseline.",
        gaps_limitations="AD not forecasting; different evaluation goals.",
        maps_to_our_gap="Module 3 foundation.",
        relevance_to_fyp="Medium (Module 3)",
        related_module="Module 3",
        our_response="Cite as deep AD foundation for later modules.",
    ),
]

GAP_LEGEND = [
    dict(
        gap_id="G1",
        gap_title="Limited generalization / per-entity models",
        description="Many approaches need separate models per host/VM/container and struggle with unseen workloads.",
    ),
    dict(
        gap_id="G2",
        gap_title="Statistical XOR deep learning",
        description="Few works rigorously combine interpretable statistical forecasting with nonlinear DL for containers.",
    ),
    dict(
        gap_id="G3",
        gap_title="Accuracy over deployability",
        description="Limited attention to scalability, ops cost, and deployment practicality.",
    ),
    dict(
        gap_id="G4",
        gap_title="Clean regular telemetry assumption",
        description="Often assume complete regularly sampled data; weak on missing/irregular monitoring.",
    ),
    dict(
        gap_id="G5",
        gap_title="Offline retraining / weak adaptive lifecycle",
        description="Limited principled post-deployment drift monitoring and retrain/deploy governance for forecasters.",
    ),
]

CONTRIB = [
    dict(
        contribution_id="C1",
        contribution="Hybrid Prophet + GRU residual forecasting for long-term container CPU",
        addresses_gaps="G2, G1",
        evidence="baseline_reference_2026-07-14; Day-1 MAE ≈ 1.746",
    ),
    dict(
        contribution_id="C2",
        contribution="Global GRU multi-container baseline",
        addresses_gaps="G1, G3",
        evidence="global_gru_baseline_2026-07-17_121748; Day-1 MAE ≈ 1.924",
    ),
    dict(
        contribution_id="C3",
        contribution="Locked Hybrid vs Global comparison protocol",
        addresses_gaps="G1, G2, G3",
        evidence="hybrid_vs_global_2026-07-17_095147; Hybrid better 67/99 by MAE",
    ),
    dict(
        contribution_id="C4",
        contribution="Peak-aware learning investigation (honest null/mixed)",
        addresses_gaps="G2, peak-error focus",
        evidence="peak_aware experiments; overall not a clear win",
    ),
    dict(
        contribution_id="C5",
        contribution="Residual diagnostics + LFHE (MSE vs DA-MSE)",
        addresses_gaps="objective vs architecture for variance collapse",
        evidence="LFHE inconclusive_I1; dispersion up, correlation not",
    ),
    dict(
        contribution_id="C6",
        contribution="AFMLF adaptive forecast model lifecycle",
        addresses_gaps="G5, G3",
        evidence="AFMLF official run; keep_current_model; ΔMAE ≈ -0.0028",
    ),
    dict(
        contribution_id="C7",
        contribution="Reproducible Alibaba preprocessing / data-quality analysis",
        addresses_gaps="G4",
        evidence="preprocessing pipeline; memory vs CPU suitability notebook",
    ),
]

CLOSEST = [
    dict(priority=1, ref_no=11, why_closest="Closest Prophet+DL hybrid for Kubernetes/containers", examiner_question="How is yours different from Prophet+LSTM K8s?"),
    dict(priority=2, ref_no=15, why_closest="Major cloud CPU hybrid DL baseline (VM)", examiner_question="How different from BHyPreC?"),
    dict(priority=3, ref_no=7, why_closest="Strong BiGRU cloud host-load forecaster", examiner_question="Why not wavelet+BiGRU only?"),
    dict(priority=4, ref_no=25, why_closest="Cloud workload + transfer/uncertainty", examiner_question="Do you handle generalization/uncertainty?"),
    dict(priority=5, ref_no=27, why_closest="Prophet foundation of Hybrid", examiner_question="Why Prophet?"),
    dict(priority=6, ref_no=36, why_closest="Modern drift-aware forecasting (OneNet)", examiner_question="Is AFMLF just OneNet?"),
    dict(priority=7, ref_no=24, why_closest="Drift detection + forecast adaptation", examiner_question="How is AFMLF different from PELT adaptation?"),
    dict(priority=8, ref_no=9, why_closest="Concept drift survey foundation", examiner_question="What is concept drift literature?"),
    dict(priority=9, ref_no=18, why_closest="Hybrid traditional+DL paradigm", examiner_question="Why hybrid statistical+DL?"),
    dict(priority=10, ref_no=33, why_closest="Residual decomposition forecasting", examiner_question="Why residual hybrid?"),
    dict(priority=11, ref_no=12, why_closest="Loss design for decomposition hybrids", examiner_question="Why LFHE/loss experiments?"),
    dict(priority=12, ref_no=34, why_closest="RNN/LSTM/GRU comparison", examiner_question="Why GRU not LSTM?"),
    dict(priority=13, ref_no=20, why_closest="Autoscaling SLO/cost systems motivation", examiner_question="Why forecasting for deployment?"),
    dict(priority=14, ref_no=30, why_closest="Container resource extremes", examiner_question="Do you handle peaks/extremes?"),
    dict(priority=15, ref_no=21, why_closest="Dynamic cloud workload prediction", examiner_question="Is dynamic ML already solved?"),
]


def _style_sheet(ws, df: pd.DataFrame, widths: list[int] | None = None, freeze: str = "B2") -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    alt = PatternFill("solid", fgColor="F2F2F2")
    high = PatternFill("solid", fgColor="E2EFDA")
    vhigh = PatternFill("solid", fgColor="C6EFCE")
    med = PatternFill("solid", fgColor="FFF2CC")
    low = PatternFill("solid", fgColor="FCE4D6")

    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for excel_row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in excel_row:
            cell.alignment = wrap
            cell.border = thin
        if excel_row[0].row % 2 == 0:
            for cell in excel_row:
                cell.fill = alt

    if "relevance_to_fyp" in df.columns:
        col_idx = list(df.columns).index("relevance_to_fyp") + 1
        for r in range(2, ws.max_row + 1):
            val = str(ws.cell(r, col_idx).value or "")
            fill = None
            if "Very High" in val:
                fill = vhigh
            elif val.startswith("High"):
                fill = high
            elif "Medium" in val:
                fill = med
            elif "Low" in val:
                fill = low
            if fill is not None:
                ws.cell(r, col_idx).fill = fill

    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    if widths:
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
    else:
        for i, col in enumerate(df.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(40, max(12, len(col) + 2))
    ws.row_dimensions[1].height = 30
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 55


def main() -> None:
    raw_lines = [
        line.strip()
        for line in (REPO / "reference.txt").read_text().splitlines()
        if line.strip()
    ]
    citation_by_ref: dict[int, str] = {}
    for i, line in enumerate(raw_lines, 1):
        match = re.match(r"^\[(\d+)\]\s*(.*)$", line)
        if match:
            ref_no = int(match.group(1))
            cite = match.group(2).strip()
        else:
            ref_no = i
            cite = re.sub(r"^\[\d+\]\s*", "", line).strip()
        citation_by_ref[ref_no] = cite

    df = pd.DataFrame(ROWS)
    df["full_ieee_citation"] = df["ref_no"].map(citation_by_ref).fillna("")
    # If [1] missing in reference.txt, use constructed citation for ref 1
    if not df.loc[df["ref_no"] == 1, "full_ieee_citation"].iloc[0] and raw_lines:
        df.loc[df["ref_no"] == 1, "full_ieee_citation"] = raw_lines[0]

    main_cols = [
        "ref_no",
        "authors",
        "year",
        "title",
        "venue",
        "doi_or_id",
        "category",
        "method_summary",
        "key_contribution",
        "gaps_limitations",
        "maps_to_our_gap",
        "relevance_to_fyp",
        "related_module",
        "our_response",
        "full_ieee_citation",
    ]
    df = df[main_cols]

    wb = Workbook()

    ws0 = wb.active
    ws0.title = "README_How_to_use"
    readme_lines = [
        ["Literature Review Reference Matrix — DracaSys FYP"],
        ["Source", "reference.txt (37 papers)"],
        [
            "Purpose",
            "Viva/literature answers: existing work, gaps per paper, mapping to our gaps/contributions",
        ],
        [],
        ["Sheets"],
        [
            "1. Literature_Matrix",
            "One row per paper: contribution, gaps, relevance, our response, full citation",
        ],
        ["2. Our_Research_Gaps", "G1–G5 from project literature review"],
        ["3. Our_Contributions", "C1–C7 what we claim + evidence"],
        ["4. Closest_Papers_Viva", "Highest-priority papers examiners may ask about"],
        ["5. Category_Summary", "Counts by category / module"],
        [],
        ["How to answer viva"],
        ["Step 1", "Filter Literature_Matrix by Very High/High or open Closest_Papers_Viva"],
        ["Step 2", "Say: Existing work does X (key_contribution)"],
        ["Step 3", "Say: Gap/limitation is Y (gaps_limitations)"],
        ["Step 4", "Say: Our work addresses this by Z (our_response)"],
        [],
        [
            "Important",
            "Do not claim we invented Prophet/GRU/Page-Hinkley; claim integrated Module-1 methodology + lifecycle",
        ],
        [
            "Module note",
            "Many anomaly/security papers support Modules 3–4 vision, not Module 1 forecasting novelty",
        ],
    ]
    for line in readme_lines:
        ws0.append(line)
    ws0["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws0.column_dimensions["A"].width = 18
    ws0.column_dimensions["B"].width = 110

    ws1 = wb.create_sheet("Literature_Matrix")
    _style_sheet(
        ws1,
        df,
        widths=[8, 22, 8, 40, 28, 28, 28, 36, 36, 40, 28, 16, 16, 40, 50],
    )

    ws2 = wb.create_sheet("Our_Research_Gaps")
    _style_sheet(ws2, pd.DataFrame(GAP_LEGEND), widths=[10, 36, 80], freeze="A2")
    for r in range(2, ws2.max_row + 1):
        ws2.row_dimensions[r].height = 40

    ws3 = wb.create_sheet("Our_Contributions")
    _style_sheet(ws3, pd.DataFrame(CONTRIB), widths=[12, 60, 28, 50], freeze="A2")
    for r in range(2, ws3.max_row + 1):
        ws3.row_dimensions[r].height = 40

    ws4 = wb.create_sheet("Closest_Papers_Viva")
    closest_df = pd.DataFrame(CLOSEST).merge(
        df[["ref_no", "authors", "year", "title", "relevance_to_fyp", "our_response"]],
        on="ref_no",
        how="left",
    )
    closest_df = closest_df[
        [
            "priority",
            "ref_no",
            "authors",
            "year",
            "title",
            "why_closest",
            "examiner_question",
            "relevance_to_fyp",
            "our_response",
        ]
    ]
    _style_sheet(ws4, closest_df, widths=[10, 8, 22, 8, 40, 40, 40, 14, 50], freeze="A2")

    ws5 = wb.create_sheet("Category_Summary")
    cat = (
        df.groupby("category")
        .agg(
            paper_count=("ref_no", "count"),
            example_refs=("ref_no", lambda s: ", ".join(map(str, sorted(s)))),
        )
        .reset_index()
        .sort_values("paper_count", ascending=False)
    )
    mod = (
        df.groupby("related_module")
        .agg(
            paper_count=("ref_no", "count"),
            example_refs=("ref_no", lambda s: ", ".join(map(str, sorted(s)))),
        )
        .reset_index()
    )
    rel = df.groupby("relevance_to_fyp").agg(paper_count=("ref_no", "count")).reset_index()

    ws5.append(["Summary by Category"])
    ws5.append(list(cat.columns))
    for row in cat.itertuples(index=False):
        ws5.append(list(row))
    ws5.append([])
    ws5.append(["Summary by Related Module"])
    ws5.append(list(mod.columns))
    for row in mod.itertuples(index=False):
        ws5.append(list(row))
    ws5.append([])
    ws5.append(["Summary by Relevance"])
    ws5.append(list(rel.columns))
    for row in rel.itertuples(index=False):
        ws5.append(list(row))
    ws5["A1"].font = Font(bold=True, size=12, color="1F4E79")
    ws5.column_dimensions["A"].width = 45
    ws5.column_dimensions["B"].width = 14
    ws5.column_dimensions["C"].width = 40

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Sheets: {wb.sheetnames}")
    print(f"Literature rows: {len(df)}")


if __name__ == "__main__":
    main()
